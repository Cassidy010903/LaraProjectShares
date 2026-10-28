import os
import re
import openpyxl

import blocks
import crosscheck
import engine
import invoice
import summary
import splitout
import writer
from engine import MONTHS, TOTALS_ROWS, ENTITY_NAMES


def guess_month(path):
    name = os.path.basename(path).lower()
    month = next((i + 1 for i, m in enumerate(MONTHS) if m.lower() in name), None)
    year_match = re.search(r"20\d\d", name)
    return month, int(year_match.group()) if year_match else None


def load_workbooks(book_path):
    wb_f = openpyxl.load_workbook(book_path)
    wb_v = openpyxl.load_workbook(book_path, data_only=True)
    return wb_f, wb_v


def check(list_path, book_path, month, year):
    res = engine.read_user_list(list_path, month, year)
    if res.errors:
        return res, None, None, None, res.errors, []
    wb_f, wb_v = load_workbooks(book_path)
    required = ("MS365 Totals", "MS365 Fees", "2026")
    missing = [s for s in required if s not in wb_f.sheetnames]
    if missing:
        return res, None, None, None, [f"The workbook does not have the sheet(s): {', '.join(missing)}."], []
    plan = engine.build_plan(res, wb_f, wb_v)
    errors = list(plan.errors)
    warnings = list(res.warnings) + list(plan.warnings)
    if not errors:
        warnings += engine.comparison_warnings(res, plan)
    return res, plan, wb_f, wb_v, errors, warnings


def output_path(book_path, month, year):
    folder = os.path.dirname(book_path)
    stem = os.path.splitext(os.path.basename(book_path))[0]
    month_pattern = "|".join(MONTHS)
    while True:
        stripped = re.sub(rf" - ({month_pattern}) \d{{4}}( \(\d+\))?$", "", stem)
        if stripped == stem:
            break
        stem = stripped
    base = os.path.join(folder, f"{stem} - {MONTHS[month - 1]} {year}")
    path = base + ".xlsx"
    n = 2
    while os.path.exists(path):
        path = f"{base} ({n}).xlsx"
        n += 1
    return path


def unique(items):
    seen = set()
    out = []
    for item in items:
        if item not in seen:
            seen.add(item)
            out.append(item)
    return out


def report_text(res, plan, data, splits, errors, warnings, invoice_notes):
    g = res.groups
    out = []
    out.append(f"{MONTHS[res.month - 1]} {res.year}: {data.users_read} users read, {int(res.seats_total)} Microsoft 365 seats, {len(data.order)} branches.")
    out.append("")
    out += data.notes + ([""] if data.notes else [])
    out.append("Microsoft 365 seats per entity (as used on '2. Beperk Verhalings'):")
    for ent in TOTALS_ROWS:
        now = int(res.entity_seats.get(ent, 0))
        before = plan.previous["entities"].get(ent) if plan is not None and plan.previous.get("entities") else None
        out.append(f"  {ENTITY_NAMES[ent]:<34}{now:>5}" + (f"   (last month {int(before)})" if before is not None else ""))
    out.append("")
    out.append("Adobe / Read AI / ChatGPT users (VS / SMA / BPK):")
    for label, key in (("Acrobat Pro", "pro"), ("Acrobat Standard DC", "std"), ("Creative Cloud", "cc"), ("Read AI", "read"), ("ChatGPT", "gpt")):
        total = sum(len(g[x][key]) for x in g)
        out.append(f"  {label:<22}{total:>3}   ({len(g['VS'][key])} / {len(g['SMA'][key])} / {len(g['BPK'][key])})")
    out.append("")
    out.append("By business unit and branch (all subscriptions):")
    out.append("")
    out.append(summary.branch_table_text(data))
    out.append(f"(Shared Inbox Active: {data.product_total('shared inbox')})")
    if splits:
        out.append("")
        out.append("INVOICE SPLIT (invoice total x branch users / total users of that subscription)")
        out.append("")
        out.append(splitout.split_summary_text(splits))
        out.append("")
        out.append("Incl VAT per business unit and branch (Excl VAT and VAT are in the Excel file):")
        out.append("")
        out.append(splitout.split_table_text(data, splits))
    if data.missing_columns:
        out.append("")
        out.append("Not in this user list (counted as zero): " + ", ".join(data.missing_columns))
    if errors:
        out.append("")
        out.append("PROBLEMS - the workbook cannot be created until these are fixed:")
        out += [f"  - {e}" for e in errors]
    if warnings:
        out.append("")
        out.append("PLEASE CHECK:")
        out += [f"  - {w}" for w in warnings]
    if invoice_notes:
        out.append("")
        out.append("INVOICES - PLEASE CHECK:")
        out += [f"  - {n}" for n in invoice_notes]
    return "\n".join(out)


class RunResult:
    def __init__(self):
        self.out_path = None
        self.errors = []
        self.warnings = []
        self.invoice_notes = []
        self.notes = []
        self.title = ""
        self.stats = {}
        self.entities = []
        self.others = []
        self.branch_heads = []
        self.branch_rows = []
        self.split_blocks = []
        self.split_heads = []
        self.split_rows = []
        self.text = ""


def run_full(list_path, book_path, month, year, invoices, out_path=None):
    result = RunResult()
    result.title = f"{MONTHS[month - 1]} {year}"
    res, plan, wb_f, wb_v, errors, warnings = check(list_path, book_path, month, year)
    data = summary.read_user_data(list_path)
    errors = unique(list(errors) + list(data.errors))
    warnings = unique(list(warnings) + list(data.warnings))
    if res.errors or plan is None:
        result.errors = errors
        result.text = "PROBLEMS:\n" + "\n".join(f"  - {e}" for e in errors)
        return result
    for inv in invoices or []:
        inv.counted = invoice.in_month(inv, month, year)
    used = [inv for inv in invoices or [] if inv.counted]
    splits = invoice.build_splits(data, used) if used else None
    split_notes = [s.note for s in splits or [] if s.note]
    notes = invoice.invoice_warnings(invoices, month, year) if invoices else None
    figures = None
    duplicates = crosscheck.duplicate_rows(wb_v, month, year)
    if invoices:
        notes = list(notes) + crosscheck.invoice_checks(wb_v, used, data, month, year, {d["row"] for d in duplicates})
        figures, figure_notes = blocks.prepare(used)
        notes += figure_notes + split_notes
        for inv in used:
            notes += inv.assumptions
    excluded_excl = sum(d["excl"] or 0 for d in duplicates)
    duplicate_notes = crosscheck.duplicate_messages(duplicates, month, year)
    formula = crosscheck.formula_notes(wb_f, wb_v, res, plan, month, year, excluded_excl)
    warnings = unique(list(warnings) + duplicate_notes + list(formula))
    result.errors = errors
    result.warnings = warnings
    result.invoice_notes = list(notes or [])
    result.notes = list(data.notes)
    result.stats = {"users": data.users_read, "seats": int(res.seats_total), "branches": len(data.order),
                    "invoices": len(used), "invoices_total": len(invoices or [])}
    for ent in TOTALS_ROWS:
        before = plan.previous["entities"].get(ent) if plan.previous.get("entities") else None
        result.entities.append((ENTITY_NAMES[ent], int(res.entity_seats.get(ent, 0)), None if before is None else int(before)))
    g = res.groups
    for label, key in (("Acrobat Pro", "pro"), ("Acrobat Standard DC", "std"), ("Creative Cloud", "cc"), ("Read AI", "read"), ("ChatGPT", "gpt")):
        result.others.append((label, sum(len(g[x][key]) for x in g), len(g["VS"][key]), len(g["SMA"][key]), len(g["BPK"][key])))
    result.branch_heads, result.branch_rows = summary.branch_table_data(data)
    if splits:
        result.split_blocks = splitout.split_blocks(splits)
        result.split_heads, result.split_rows = splitout.split_table_data(data, splits)
    result.text = report_text(res, plan, data, splits, errors, warnings, notes)
    if errors:
        return result
    result.out_path = out_path or output_path(book_path, month, year)
    writer.write_workbook(res, plan, wb_f, result.out_path, unique(list(warnings) + list(notes or [])), data, invoices or None, splits, figures, duplicates)
    return result


def run(list_path, book_path, month, year, invoices, out_path=None):
    result = run_full(list_path, book_path, month, year, invoices, out_path)
    return result.text, result.out_path

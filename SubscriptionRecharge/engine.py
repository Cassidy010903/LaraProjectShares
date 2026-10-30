import re
import json
import os
import sys
import datetime
import collections
import openpyxl

MONTHS = ["January", "February", "March", "April", "May", "June", "July",
          "August", "September", "October", "November", "December"]

MS_SEAT_COLUMNS = ["Business Standard", "Business Premium", "F3", "CoPilot",
                   "Microsoft Power Bi Pro", "Microsoft Power Bi Premium", "Teams Premium"]
PRO = "Adobe Acrobat Pro"
STD = "Adobe Acrobat Standard DC"
CC = "Adobe Creative Suite"
READ = "Read AI"
GPT = ["Chat GPT", "Chat GPT Premium"]
TRACKED = MS_SEAT_COLUMNS + [PRO, STD, CC, READ] + GPT

TOTALS_ROWS = {"SMA": 5, "KAAP": 6, "DBN": 7, "ONLINE": 8, "GWP": 9, "LTD": 10,
               "JHB": 11, "WENPRO": 12, "DWF PTA": 13, "Zire DC": 16, "VS": 17}
ENTITY_NAMES = {"SMA": "SMA (Subtropico Market Agents)", "KAAP": "KAAP (Cape Town)", "DBN": "DBN (Durban)",
                "ONLINE": "ONLINE", "GWP": "GWP", "LTD": "LTD / BPK (Subtropico Limited)", "JHB": "JHB (Johannesburg)",
                "WENPRO": "WENPRO", "DWF PTA": "DWF PTA (DW Fresh)", "Zire DC": "Zire DC", "VS": "VS (Vleissentraal)"}
SMA_TABLE_FIRST, SMA_TABLE_LAST, SMA_TABLE_HEAD, SMA_TABLE_SUM = 29, 39, 28, 40
VS_TABLE_FIRST, VS_TABLE_LAST, VS_TABLE_HEAD, VS_TABLE_SUM = 46, 60, 45, 62
VAT = 0.15

DEFAULT_MAP = {
    "Subtropico": {
        "smabloemfontein": ("SMA", "BFN"), "bloemfontein": ("SMA", "BFN"),
        "smaeastlondon": ("SMA", "EAST LONDON"), "eastlondon": ("SMA", "EAST LONDON"),
        "smaheadoffice": ("SMA", "HK"), "smaheadofficepretoria": ("SMA", "HK"),
        "smaklerksdorp": ("SMA", "KDRP"), "klerksdorp": ("SMA", "KDRP"),
        "smakimberley": ("SMA", "KIM"), "kimberley": ("SMA", "KIM"),
        "smapretoria": ("SMA", "PTA"), "pretoria": ("SMA", "PTA"),
        "smasprings": ("SMA", "SPR"), "springs": ("SMA", "SPR"),
        "smavereeniging": ("SMA", "VER"), "vereeniging": ("SMA", "VER"),
        "smawelkom": ("SMA", "WEL"), "welkom": ("SMA", "WEL"),
        "smawitbank": ("SMA", "WIT"), "witbank": ("SMA", "WIT"),
        "smapietermaritzburg": ("SMA", "PMB"), "pietermaritzburg": ("SMA", "PMB"),
        "smacapetown": ("KAAP", None), "capetown": ("KAAP", None),
        "smadurban": ("DBN", None), "durban": ("DBN", None),
        "smadwfresh": ("DWF PTA", None), "dwfresh": ("DWF PTA", None),
        "smajohannesburg": ("JHB", None), "johannesburg": ("JHB", None), "subjohannesburg": ("JHB", None),
        "smajohannesburgwenpro": ("WENPRO", None), "wenpro": ("WENPRO", None),
        "limited": ("LTD", None), "headoffice": ("LTD", None),
        "online": ("ONLINE", None), "gwp": ("GWP", None),
    },
    "Vleissentraal": {
        "bethlehem": ("VS", "Bethlehem"), "bethlehemkroonstad": ("VS", "BHM - Kroonstad"),
        "bloemfontein": ("VS", "Bloemfontein"), "bosveld": ("VS", "Bosveld"),
        "ermelo": ("VS", "Ermelo"), "frankfort": ("VS", "Frankfort"),
        "headoffice": ("VS", "Head Office"), "kimberley": ("VS", "Kimberley"),
        "klerksdorp": ("VS", "Klerksdorp"), "kzn": ("VS", "KZN"), "ladysmith": ("VS", "KZN"),
        "kaap": ("VS", "Kaap"), "noordkaap": ("VS", "Northern Cape"),
        "suidkaap": ("VS", "Suid-Kaap"), "portelizabeth": ("VS", "PE"),
        "property": ("VS", "Property"),
    },
}

GROUP_OF_ENTITY = {"VS": "VS", "LTD": "BPK", "Zire DC": "ZIRE"}
GROUP_ENTITY = {"SMA": "SMA", "VS": "VS", "BPK": "LTD", "ZIRE": "Zire DC"}


def app_dir():
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))


def alnum(text):
    return re.sub(r"[^a-z0-9]", "", str(text).lower())


def load_mapping():
    mapping = {bu: dict(v) for bu, v in DEFAULT_MAP.items()}
    path = os.path.join(app_dir(), "branch_map.json")
    if os.path.exists(path):
        with open(path, encoding="utf-8") as f:
            extra = json.load(f)
        for key, value in extra.items():
            bu, branch = key.split("|", 1)
            mapping.setdefault(bu.strip(), {})[alnum(branch)] = (value[0], value[1] if len(value) > 1 else None)
    return mapping


def num(value):
    if isinstance(value, bool):
        return 0
    if isinstance(value, (int, float)):
        return value
    return 0


def clean_header(value):
    return re.sub(r"\s+", " ", str(value)).strip().lower() if value is not None else ""


def first_name(label, firstnames, upn):
    if firstnames:
        return str(firstnames).strip()
    text = str(label or "").strip()
    return text.split(" ")[0] if text else str(upn or "")


def join_names(names):
    if len(names) > 1:
        return ", ".join(names[:-1]) + " & " + names[-1]
    return names[0] if names else ""


class Result:
    def __init__(self, month, year):
        self.errors = []
        self.warnings = []
        self.info = []
        self.month = month
        self.year = year
        self.seats_total = 0
        self.entity_seats = collections.Counter()
        self.sma_table = collections.Counter()
        self.vs_table = collections.Counter()
        self.licence_totals = collections.Counter()
        self.groups = {g: {"pro": [], "std": [], "cc": [], "read": [], "gpt": []} for g in ("SMA", "VS", "BPK", "ZIRE")}
        self.users_read = 0


def read_user_list(path, month, year):
    res = Result(month, year)
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb["Users"] if "Users" in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    header_index = None
    for i, r in enumerate(rows[:15]):
        if r and clean_header(r[0]) == "label":
            header_index = i
            break
    if header_index is None:
        res.errors.append("Could not find the header row (a row starting with 'Label') on the Users sheet.")
        return res
    cols = {}
    for j, h in enumerate(rows[header_index]):
        key = clean_header(h)
        if key:
            cols.setdefault(key, j)
    missing = [n for n in ("upn", "business unit", "branch", "business standard", "f3") if n not in cols]
    if missing:
        res.errors.append("The user list is missing required columns: " + ", ".join(missing))
        return res
    for name in TRACKED:
        if clean_header(name) not in cols:
            res.info.append(f"Column '{name}' is not in this user list (treated as zero).")

    def cell(r, name):
        j = cols.get(clean_header(name))
        return r[j] if j is not None and j < len(r) else None

    mapping = load_mapping()
    unknown = collections.OrderedDict()
    seen_upn = collections.Counter()
    for r in rows[header_index + 1:]:
        if r is None:
            continue
        label, upn = cell(r, "label"), cell(r, "upn")
        if label is None and upn is None:
            continue
        lic = {c: num(cell(r, c)) for c in TRACKED}
        if not any(lic.values()):
            continue
        res.users_read += 1
        bu = str(cell(r, "business unit") or "").strip()
        branch = str(cell(r, "branch") or "").strip()
        who = str(label or upn)
        if upn:
            seen_upn[str(upn).strip().lower()] += 1
        if bu == "Zire DC":
            target = ("Zire DC", None)
        else:
            target = mapping.get(bu, {}).get(alnum(branch))
        if target is None:
            unknown.setdefault((bu, branch), []).append(who)
            continue
        entity, table_label = target
        seats = sum(lic[c] for c in MS_SEAT_COLUMNS)
        for c in TRACKED:
            res.licence_totals[c] += lic[c]
        res.seats_total += seats
        res.entity_seats[entity] += seats
        if seats and entity == "SMA" and table_label:
            res.sma_table[table_label] += seats
        if seats and entity == "VS" and table_label:
            res.vs_table[table_label] += seats
        g = res.groups[GROUP_OF_ENTITY.get(entity, "SMA")]
        name = first_name(label, cell(r, "firstnames"), upn)
        g["pro"] += [name] * int(lic[PRO])
        g["std"] += [name] * int(lic[STD])
        g["cc"] += [name] * int(lic[CC])
        g["read"] += [name] * int(lic[READ])
        g["gpt"] += [name] * int(sum(lic[c] for c in GPT))
        if bu == "Subtropico" and alnum(branch) == "headoffice":
            res.info.append(f"Subtropico 'Head Office' user {who} is counted under LTD (Subtropico Limited).")
    for (bu, branch), names in unknown.items():
        res.errors.append(f"Unknown Business Unit / Branch '{bu}' / '{branch or '(blank)'}' on {len(names)} licensed user(s): "
                          + ", ".join(names[:6]) + (" ..." if len(names) > 6 else "")
                          + ". Fix it in the user list, or add it to branch_map.json next to the program.")
    dupes = [u for u, n in seen_upn.items() if n > 1]
    if dupes:
        res.warnings.append("Duplicate accounts (the same UPN appears more than once, so seats are counted twice): " + ", ".join(dupes[:6]))
    zg = res.groups["ZIRE"]
    if any(zg[k] for k in zg):
        res.warnings.append("Zire DC has Adobe / Read AI / ChatGPT licences, but the workbook blocks have no Zire section, so they are only shown on the Sage sheet.")
    return res


def find_col_by_date(ws, row, month, year):
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row, c).value
        if isinstance(v, datetime.datetime) and v.month == month and v.year == year:
            return c
    return None


def find_col_by_text(ws, row, text):
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row, c).value
        if isinstance(v, str) and v.strip().lower() == text.lower():
            return c
    return None


def find_block(ws, kind, month):
    abbr = MONTHS[month - 1][:3].lower()
    for c in range(1, ws.max_column + 1):
        v = ws.cell(1, c).value
        if isinstance(v, str):
            t = v.strip().lower()
            if t.startswith(kind) and re.search(r"\b" + abbr, t):
                return c
    return None


def month_label(month, year):
    return f"{MONTHS[month - 1][:3]} {str(year)[2:]}"


def locate(wb_f, month, year):
    tot = wb_f["MS365 Totals"]
    sheet = wb_f["2026"]
    return {
        "count_col": find_col_by_date(tot, 4, month, year),
        "sma_col": find_col_by_text(tot, SMA_TABLE_HEAD, month_label(month, year)),
        "vs_col": find_col_by_text(tot, VS_TABLE_HEAD, month_label(month, year)),
        "adobe": find_block(sheet, "adobe", month),
        "read": find_block(sheet, "read", month),
        "gpt": find_block(sheet, "openai", month),
    }


def cached(ws, row, col):
    v = ws.cell(row, col).value
    return v if isinstance(v, (int, float)) else 0


def table_labels(ws, first, last):
    labels = {}
    for r in range(first, last + 1):
        v = ws.cell(r, 8).value
        if isinstance(v, str) and v.strip():
            labels[v.strip().upper()] = r
    return labels


class Plan:
    def __init__(self):
        self.errors = []
        self.warnings = []
        self.pos = None
        self.sma_labels = {}
        self.vs_labels = {}
        self.previous = {}


def build_plan(res, wb_f, wb_v):
    plan = Plan()
    m, y = res.month, res.year
    pos = locate(wb_f, m, y)
    plan.pos = pos
    names = {"count_col": "month column on 'MS365 Totals'", "sma_col": "SMA branch table column on 'MS365 Totals'",
             "vs_col": "Vleissentraal branch table column on 'MS365 Totals'",
             "adobe": "Adobe block on '2026'", "read": "Read AI block on '2026'", "gpt": "OpenAI block on '2026'"}
    optional = ("adobe", "read", "gpt")
    for k, label in names.items():
        if pos.get(k) is None and k not in optional:
            plan.errors.append(f"Could not find the {label} for {MONTHS[m - 1]} {y} in the workbook.")
    missing_blocks = [names[k] for k in optional if pos.get(k) is None]
    if missing_blocks:
        plan.warnings.append(f"The {', '.join(missing_blocks)} for {MONTHS[m - 1]} {y} do not exist in the workbook yet, so those user counts were not filled in there. "
                             "The counts are in the Summary sheets.")
    if plan.errors:
        return plan
    tot_f = wb_f["MS365 Totals"]
    plan.sma_labels = table_labels(tot_f, SMA_TABLE_FIRST, SMA_TABLE_LAST)
    plan.vs_labels = table_labels(tot_f, VS_TABLE_FIRST, VS_TABLE_LAST)
    for label in res.sma_table:
        if label.upper() not in plan.sma_labels:
            plan.errors.append(f"Branch '{label}' has no row in the SMA branch table on 'MS365 Totals'.")
    for label in res.vs_table:
        if label.upper() not in plan.vs_labels:
            plan.errors.append(f"Branch '{label}' has no row in the Vleissentraal branch table on 'MS365 Totals'.")
    if res.seats_total <= 0:
        plan.errors.append("No Microsoft seats were found in the user list.")
    plan.previous = previous_counts(wb_v, m, y)
    return plan


def previous_counts(wb_v, month, year):
    out = {"entities": {}, "sma": {}, "vs": {}}
    if month == 1:
        return out
    tot = wb_v["MS365 Totals"]
    col = find_col_by_date(tot, 4, month - 1, year)
    if col:
        for ent, row in TOTALS_ROWS.items():
            out["entities"][ent] = cached(tot, row, col)
    c1 = find_col_by_text(tot, SMA_TABLE_HEAD, month_label(month - 1, year))
    c2 = find_col_by_text(tot, VS_TABLE_HEAD, month_label(month - 1, year))
    for first, last, c, key in ((SMA_TABLE_FIRST, SMA_TABLE_LAST, c1, "sma"), (VS_TABLE_FIRST, VS_TABLE_LAST, c2, "vs")):
        if c:
            for r in range(first, last + 1):
                label = tot.cell(r, 8).value
                if label:
                    out[key][str(label).strip().upper()] = cached(tot, r, c)
    for entity, key in (("SMA", "sma"), ("VS", "vs")):
        if out["entities"] and not out["entities"].get(entity) and out[key]:
            out["entities"][entity] = sum(out[key].values())
    return out


def comparison_warnings(res, plan):
    out = []
    prev = plan.previous
    if not prev.get("entities"):
        return out
    for key, now in (("sma", res.sma_table), ("vs", res.vs_table)):
        labels = plan.sma_labels if key == "sma" else plan.vs_labels
        for label in labels:
            before = prev[key].get(label, 0)
            after = now.get(next((k for k in now if k.upper() == label), ""), 0)
            if before and after < before * 0.5 and before - after >= 2:
                out.append(f"{label}: {int(before)} seats last month, {int(after)} now. Check that those users were not left off the list or given a different branch name.")
            if after - before >= 5:
                out.append(f"{label}: {int(before)} seats last month, {int(after)} now (+{int(after - before)}).")
    before_total = sum(prev["entities"].values())
    if before_total and abs(res.seats_total - before_total) / before_total > 0.1:
        out.append(f"Total seats changed from {int(before_total)} last month to {int(res.seats_total)} now.")
    return out

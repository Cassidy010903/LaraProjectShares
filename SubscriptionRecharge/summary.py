import re
import collections
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

import brand

MONTHS = ["January", "February", "March", "April", "May", "June", "July",
          "August", "September", "October", "November", "December"]

PRODUCTS = [
    ("shared inbox", "Shared Inbox", "Shared Inbox Active"),
    ("business standard", "Business Standard", "Business Standard"),
    ("business premium", "Business Premium", "Business Premium"),
    ("f3", "F3", "F3"),
    ("teams premium", "Teams Premium", "Teams Premium"),
    ("copilot", "CoPilot", "CoPilot"),
    ("microsoft power bi pro", "Microsoft Power Bi Pro", "Microsoft Power Bi Pro"),
    ("microsoft power bi premium", "Microsoft Power Bi Premium", "Microsoft Power Bi Premium"),
    ("adobe acrobat pro", "Adobe Acrobat Pro", "Adobe Acrobat Pro"),
    ("adobe acrobat standard dc", "Adobe Acrobat Standard DC", "Adobe Acrobat Standard DC"),
    ("adobe creative suite", "Adobe Creative Suite", "Adobe Creative Suite"),
    ("read ai", "Read AI", "Read AI"),
    ("chat gpt", "Chat GPT", "Chat GPT"),
    ("chat gpt premium", "Chat GPT Premium", "Chat GPT Premium"),
    ("fresh service agent", "Fresh Service Agent", "Fresh Service Agent"),
    ("canva", "Canva", "Canva"),
]
BASE_KEYS = [p[0] for p in PRODUCTS]

BASE_GROUPS = [
    ("Microsoft 365", ["business standard", "business premium", "f3", "teams premium", "copilot",
                       "microsoft power bi pro", "microsoft power bi premium"]),
    ("Adobe", ["adobe acrobat pro", "adobe acrobat standard dc", "adobe creative suite"]),
    ("Read AI", ["read ai"]),
    ("Chat GPT", ["chat gpt", "chat gpt premium"]),
    ("Fresh Service Agent", ["fresh service agent"]),
    ("Canva", ["canva"]),
]

BASE_LAYOUT = [
    ("shared inbox", None), ("business standard", None), ("business premium", None), ("f3", None),
    ("teams premium", None), ("copilot", None), ("microsoft power bi pro", None),
    ("microsoft power bi premium", None), (None, "Microsoft 365"),
    ("adobe acrobat pro", None), ("adobe acrobat standard dc", None), ("adobe creative suite", None), (None, "Adobe"),
    ("read ai", None), ("chat gpt", None), ("chat gpt premium", None), (None, "Chat GPT"),
    ("fresh service agent", None), ("canva", None),
]

NOT_SUBSCRIPTIONS = {"label", "upn", "role", "aliases", "device name", "firstnames", "surname", "contact number",
                     "deleted", "business unit", "branch", "creation date", "cancellation date", "one drive",
                     "notes", "assigment", "assignment"}

HEAD_FILL = PatternFill("solid", fgColor=brand.GREEN)
GROUP_FILL = PatternFill("solid", fgColor=brand.GREY_LIGHT)
NEW_FILL = PatternFill("solid", fgColor=brand.YELLOW)
BOLD = Font(name=brand.FONT, bold=True)


def clean(value):
    return re.sub(r"\s+", " ", str(value)).strip().lower() if value is not None else ""


def display(value):
    return re.sub(r"\s+", " ", str(value)).strip()


def parse_price(value):
    if isinstance(value, bool):
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        text = re.sub(r"[^0-9,.\-]", "", value)
        if "," in text and "." not in text:
            text = text.replace(",", ".")
        else:
            text = text.replace(",", "")
        try:
            return float(text)
        except ValueError:
            return 0.0
    return 0.0


class UserData:
    def __init__(self):
        self.errors = []
        self.warnings = []
        self.notes = []
        self.prices = {}
        self.order = []
        self.counts = collections.defaultdict(lambda: collections.Counter())
        self.names = collections.defaultdict(lambda: collections.defaultdict(list))
        self.users_read = 0
        self.missing_columns = []
        self.keys = list(BASE_KEYS)
        self.detail_name = {p[0]: p[1] for p in PRODUCTS}
        self.breakdown_name = {p[0]: p[2] for p in PRODUCTS}
        self.groups = list(BASE_GROUPS)
        self.layout = list(BASE_LAYOUT)
        self.extras = []

    def group_count(self, key, group_name):
        members = dict(self.groups)[group_name]
        return sum(self.counts[key][m] for m in members)

    def product_total(self, product):
        return sum(self.counts[k][product] for k in self.order)

    def group_total(self, group_name):
        return sum(self.group_count(k, group_name) for k in self.order)


def find_extras(rows, header_index, cols):
    import openpyxl
    extras = []
    known = set(BASE_KEYS) | NOT_SUBSCRIPTIONS
    start = cols.get("shared inbox", cols.get("deleted"))
    end = cols["business unit"]
    header = rows[header_index]
    inside = set(range(start + 1, end)) if start is not None and start < end else set()
    for j, h in enumerate(header):
        key = clean(h)
        if not key or key in known or key in {e[0] for e in extras}:
            continue
        if j in inside:
            extras.append((key, display(h), j, "between the licence columns and Business Unit"))
            continue
        if j > end:
            values = [r[j] for r in rows[header_index + 1:] if r is not None and j < len(r) and r[j] not in (None, "")]
            if values and all(isinstance(v, (int, float)) and not isinstance(v, bool) and v == 1 for v in values):
                extras.append((key, display(h), j, "after the Branch column, filled only with 1s"))
    return extras


def read_user_data(path):
    import openpyxl
    data = UserData()
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb["Users"] if "Users" in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    header_index = None
    for i, r in enumerate(rows[:15]):
        if r and clean(r[0]) == "label":
            header_index = i
            break
    if header_index is None:
        data.errors.append("Could not find the header row (a row starting with 'Label') on the Users sheet.")
        return data
    cols = {}
    for j, h in enumerate(rows[header_index]):
        key = clean(h)
        if key:
            cols.setdefault(key, j)
    for needed in ("upn", "business unit", "branch"):
        if needed not in cols:
            data.errors.append(f"The user list is missing the column '{needed}'.")
    if data.errors:
        return data
    for key, name, j, where in find_extras(rows, header_index, cols):
        cols[key] = j
        data.keys.append(key)
        data.extras.append(key)
        data.detail_name[key] = name
        data.breakdown_name[key] = name
        data.groups.append((name, [key]))
        data.layout.append((key, None))
    price_row = rows[header_index - 1] if header_index >= 1 else ()
    for key in data.keys:
        j = cols.get(key)
        if j is None:
            data.missing_columns.append(data.detail_name[key])
            data.prices[key] = 0.0
        else:
            data.prices[key] = parse_price(price_row[j]) if j < len(price_row) else 0.0
    unplaced = collections.defaultdict(list)
    odd_values = collections.Counter()
    canonical = {}
    merged = []
    for r in rows[header_index + 1:]:
        if r is None:
            continue
        label = r[cols["label"]] if "label" in cols and cols["label"] < len(r) else None
        upn = r[cols["upn"]] if cols["upn"] < len(r) else None
        if label is None and upn is None:
            continue
        data.users_read += 1
        raw_bu = r[cols["business unit"]] if cols["business unit"] < len(r) else None
        raw_branch = r[cols["branch"]] if cols["branch"] < len(r) else None
        bu = re.sub(r"\s+", " ", str(raw_bu)).strip() if raw_bu is not None else ""
        branch = re.sub(r"\s+", " ", str(raw_branch)).strip() if raw_branch is not None else ""
        name = str(label if label is not None else upn).strip()
        flags = []
        for key in data.keys:
            j = cols.get(key)
            value = r[j] if j is not None and j < len(r) else None
            if isinstance(value, (int, float)) and not isinstance(value, bool) and value == 1:
                flags.append(key)
            elif value not in (None, "", 0) and not (isinstance(value, (int, float)) and value == 0):
                odd_values[(data.detail_name[key], str(value))] += 1
        if not bu or not branch:
            if flags:
                unplaced[(bu or "(blank)", branch or "(blank)")].append(name)
            continue
        same = (bu.lower(), re.sub(r"[^a-z0-9]", "", branch.lower()))
        if same not in canonical:
            canonical[same] = (bu, branch)
        elif canonical[same][1] != branch and (canonical[same][1], branch) not in merged:
            merged.append((canonical[same][1], branch))
            data.notes.append(f"Merged the branch spelled '{branch}' into '{canonical[same][1]}' ({canonical[same][0]}): same business unit, same branch spelled differently.")
        key = canonical[same]
        if key not in data.order:
            data.order.append(key)
        for product in flags:
            data.counts[key][product] += 1
            data.names[key][product].append(name)
    for key in data.extras:
        total = data.product_total(key)
        price = data.prices.get(key, 0.0)
        data.notes.append(f"NEW SUBSCRIPTION FOUND: '{data.detail_name[key]}' with {total} user(s)"
                          + (f" at a price of {price:g} each." if price else " (no price in the price row, so its value shows as 0)."))
        data.notes.append("   It is included in all the summary sheets. It is not filled in on '2. Beperk Verhalings' yet.")
    for (bu, branch), names in unplaced.items():
        data.warnings.append(f"{len(names)} user(s) with licences have Business Unit '{bu}' and Branch '{branch}', so they are not in any branch row: "
                             + ", ".join(names[:6]) + (" ..." if len(names) > 6 else ""))
    for (product, value), n in odd_values.items():
        data.warnings.append(f"{n} cell(s) in the '{product}' column contain '{value}' instead of 1, so they were not counted.")
    return data


def style_header(cell, fill=HEAD_FILL):
    dark_text = fill is not HEAD_FILL
    cell.font = Font(name=brand.FONT, bold=True, color=brand.GREY_DARK if dark_text else brand.WHITE)
    cell.fill = fill
    cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")


def sheet_name(base, month, year):
    return f"{base} {MONTHS[month - 1][:3]} {year}"


def new_sheet(wb, name):
    if name in wb.sheetnames:
        del wb[name]
    return wb.create_sheet(name)


def write_summary_sheet(wb, data, title, name):
    ws = new_sheet(wb, name)
    brand.add_title(ws, f"Summary - {title}")
    for i, h in enumerate(["Business Unit", "Branch"], 1):
        style_header(ws.cell(3, i, h))
    for i, (gname, members) in enumerate(data.groups, 3):
        style_header(ws.cell(3, i, gname), NEW_FILL if members[0] in data.extras else HEAD_FILL)
    row = 4
    for key in data.order:
        ws.cell(row, 1, key[0])
        ws.cell(row, 2, key[1])
        for i, (gname, _) in enumerate(data.groups, 3):
            ws.cell(row, i, data.group_count(key, gname))
        row += 1
    ws.cell(row, 2, "Total").font = BOLD
    for i, (gname, _) in enumerate(data.groups, 3):
        ws.cell(row, i, data.group_total(gname)).font = BOLD
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 30
    for i in range(3, 3 + len(data.groups)):
        ws.column_dimensions[get_column_letter(i)].width = 16
    ws.freeze_panes = "C4"


def write_breakdown_sheet(wb, data, title, name):
    ws = new_sheet(wb, name)
    brand.add_title(ws, f"Breakdown - {title}")
    ws.cell(2, 2, "Price").font = BOLD
    for i, h in enumerate(["Business Unit", "Branch"], 1):
        style_header(ws.cell(3, i, h))
    col = 3
    spans = []
    for product, group in data.layout:
        if product:
            style_header(ws.cell(3, col, data.breakdown_name[product]), NEW_FILL if product in data.extras else HEAD_FILL)
            if data.prices.get(product):
                ws.cell(2, col, data.prices[product])
        else:
            style_header(ws.cell(3, col, group), GROUP_FILL)
        spans.append((product, group, col))
        col += 1
    row = 4
    for key in data.order:
        ws.cell(row, 1, key[0])
        ws.cell(row, 2, key[1])
        for product, group, c in spans:
            ws.cell(row, c, data.counts[key][product] if product else data.group_count(key, group))
            if group:
                ws.cell(row, c).fill = GROUP_FILL
        row += 1
    ws.cell(row, 2, "Total").font = BOLD
    for product, group, c in spans:
        ws.cell(row, c, data.product_total(product) if product else data.group_total(group)).font = BOLD
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 30
    for _, _, c in spans:
        ws.column_dimensions[get_column_letter(c)].width = 16
    ws.row_dimensions[3].height = 45
    ws.freeze_panes = "C4"


def write_detailed_sheet(wb, data, title, name):
    ws = new_sheet(wb, name)
    brand.add_title(ws, f"Detailed Breakdown - {title}")
    ws.cell(2, 2, "Price").font = BOLD
    for i, h in enumerate(["Business Unit", "Branch"], 1):
        style_header(ws.cell(3, i, h))
    col = 3
    spans = []
    for key in data.keys:
        pname = data.detail_name[key]
        fill = NEW_FILL if key in data.extras else HEAD_FILL
        if key == "shared inbox":
            style_header(ws.cell(3, col, "Shared Inbox Active"))
            style_header(ws.cell(3, col + 1, "Shared Inbox Total"))
        else:
            style_header(ws.cell(3, col, pname), fill)
            style_header(ws.cell(3, col + 1, f"{pname} Total (Rand Value)"), fill)
        style_header(ws.cell(3, col + 2, f"{pname} Users"), fill)
        if data.prices.get(key):
            ws.cell(2, col, data.prices[key])
        spans.append((key, col))
        col += 3
    row = 4
    for bkey in data.order:
        ws.cell(row, 1, bkey[0])
        ws.cell(row, 2, bkey[1])
        for key, c in spans:
            n = data.counts[bkey][key]
            ws.cell(row, c, n)
            ws.cell(row, c + 1, n * data.prices.get(key, 0.0)).number_format = "#,##0.00"
            names = ws.cell(row, c + 2, ", ".join(data.names[bkey][key]) if n else "No users")
            names.alignment = Alignment(wrap_text=True, vertical="top")
        row += 1
    ws.cell(row, 2, "Total").font = BOLD
    for key, c in spans:
        ws.cell(row, c, data.product_total(key)).font = BOLD
        total = ws.cell(row, c + 1, data.product_total(key) * data.prices.get(key, 0.0))
        total.font = BOLD
        total.number_format = "#,##0.00"
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 30
    for key, c in spans:
        ws.column_dimensions[get_column_letter(c)].width = 14
        ws.column_dimensions[get_column_letter(c + 1)].width = 18
        ws.column_dimensions[get_column_letter(c + 2)].width = 40
    ws.row_dimensions[3].height = 45
    ws.freeze_panes = "C4"


def add_summary_sheets(wb, data, month, year, invoices=None, splits=None):
    title = f"{MONTHS[month - 1]} {year}"
    write_summary_sheet(wb, data, title, sheet_name("Summary", month, year))
    write_breakdown_sheet(wb, data, title, sheet_name("Breakdown", month, year))
    write_detailed_sheet(wb, data, title, sheet_name("Detailed Breakdown", month, year))
    names = [sheet_name("Summary", month, year), sheet_name("Breakdown", month, year), sheet_name("Detailed Breakdown", month, year)]
    if invoices:
        import splitout
        split_name = sheet_name("Invoice Split", month, year)
        invoices_name = sheet_name("Invoices", month, year)
        if splits:
            splitout.write_split_sheet(wb, data, splits, title, split_name)
            names.append(split_name)
        splitout.write_invoices_sheet(wb, invoices, title, invoices_name)
        names.append(invoices_name)
    return names


def branch_table_text(data):
    heads = [g[0] if len(g[0]) <= 13 else g[0][:12] + "." for g in data.groups]
    heads = ["Fresh Svc" if h.startswith("Fresh") else h for h in heads]
    width = 30
    cw = 14
    lines = [f"{'Business Unit / Branch':<{width}}" + "".join(f"{h:>{cw}}" for h in heads)]
    lines.append("-" * (width + cw * len(heads)))
    by_bu = collections.OrderedDict()
    for key in data.order:
        by_bu.setdefault(key[0], []).append(key)
    for bu, keys in by_bu.items():
        totals = [sum(data.group_count(k, g) for k in keys) for g, _ in data.groups]
        lines.append(f"{bu:<{width}}" + "".join(f"{v:>{cw}}" for v in totals))
        for key in keys:
            vals = [data.group_count(key, g) for g, _ in data.groups]
            lines.append(f"{'   ' + key[1]:<{width}}" + "".join(f"{v:>{cw}}" for v in vals))
        lines.append("")
    totals = [data.group_total(g) for g, _ in data.groups]
    lines.append(f"{'TOTAL':<{width}}" + "".join(f"{v:>{cw}}" for v in totals))
    return "\n".join(lines)


def branch_table_data(data):
    heads = [g[0] for g in data.groups]
    by_bu = collections.OrderedDict()
    for key in data.order:
        by_bu.setdefault(key[0], []).append(key)
    rows = []
    for bu, keys in by_bu.items():
        rows.append(("bu", bu, [sum(data.group_count(k, g) for k in keys) for g, _ in data.groups]))
        for key in keys:
            rows.append(("branch", key[1], [data.group_count(key, g) for g, _ in data.groups]))
    rows.append(("total", "TOTAL", [data.group_total(g) for g, _ in data.groups]))
    return heads, rows

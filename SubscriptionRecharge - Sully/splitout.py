import collections
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

import brand
import summary

MONEY = "#,##0.00"
HEAD_FILL = PatternFill("solid", fgColor=brand.GREEN)
BU_FILL = PatternFill("solid", fgColor=brand.GREY_LIGHT)
TOTAL_FILL = PatternFill("solid", fgColor=brand.YELLOW_LIGHT)
ERROR_FILL = PatternFill("solid", fgColor="F8D7DA")


def font(bold=False, color=brand.GREY_DARK, size=10):
    return Font(name=brand.FONT, bold=bold, color=color, size=size)


def head(cell, text):
    cell.value = text
    cell.font = font(True, brand.WHITE)
    cell.fill = HEAD_FILL
    cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")


def block_title(split):
    count = len(split.invoices)
    licence = f" - {split.licence}" if split.licence and split.licence.lower() != split.subscription.lower() else ""
    return f"{split.subscription or '(no subscription)'}{licence} ({split.currency}) - {count} invoice{'s' if count != 1 else ''}"


def by_business_unit(data):
    groups = collections.OrderedDict()
    for key in data.order:
        groups.setdefault(key[0], []).append(key)
    return groups


def write_split_sheet(wb, data, splits, title, name):
    ws = summary.new_sheet(wb, name)
    brand.add_title(ws, f"Invoice Split - {title}", "F1")
    ws.cell(2, 1, "Amount per branch = amount for the licence type x (users of the branch with that licence type / all users with that licence type)").font = font(False, brand.GREY_MID)
    for i, h in enumerate(["Business Unit", "Branch"], 1):
        head(ws.cell(5, i, h), h)
    groups = by_business_unit(data)
    first_row = 6
    row_of = {}
    row = first_row
    for bu, keys in groups.items():
        row_of[("bu", bu)] = row
        row += 1
        for key in keys:
            row_of[key] = row
            row += 1
    total_row = row
    for bu, keys in groups.items():
        r = row_of[("bu", bu)]
        c = ws.cell(r, 1, bu)
        c.font = font(True)
        for col in range(1, 3 + 5 * len(splits)):
            ws.cell(r, col).fill = BU_FILL
        for key in keys:
            ws.cell(row_of[key], 2, "   " + key[1]).font = font()
    ws.cell(total_row, 2, "TOTAL").font = font(True)
    ws.cell(4, 2, "Invoice total").font = font(True)
    col = 3
    for split in splits:
        span = 5
        ws.merge_cells(start_row=3, start_column=col, end_row=3, end_column=col + span - 1)
        head(ws.cell(3, col), block_title(split))
        if split.error:
            ws.merge_cells(start_row=4, start_column=col, end_row=4, end_column=col + span - 1)
            ws.cell(4, col, split.error).font = font(True, "9C0006")
            ws.cell(4, col).fill = ERROR_FILL
            ws.cell(4, col).alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[4].height = 45
            col += span
            continue
        for i, h in enumerate(["Users", "Share", "Excl VAT", "VAT", "Incl VAT"]):
            head(ws.cell(5, col + i), h)
        ws.cell(4, col, split.total_users).font = font(True)
        ws.cell(4, col + 1, 1).number_format = "0.0%"
        for i, value in enumerate((split.excl, split.vat, split.incl)):
            c = ws.cell(4, col + 2 + i, value)
            c.number_format = MONEY
            c.font = font(True)
        for k in range(span):
            ws.cell(4, col + k).fill = TOTAL_FILL
        for bu, keys in groups.items():
            bu_users = sum(split.users[k] for k in keys)
            bu_vals = [sum(split.rows[k][j] for k in keys) for j in range(3)]
            r = row_of[("bu", bu)]
            ws.cell(r, col, bu_users).font = font(True)
            ws.cell(r, col + 1, bu_users / split.total_users).number_format = "0.0%"
            ws.cell(r, col + 1).font = font(True)
            for j in range(3):
                c = ws.cell(r, col + 2 + j, round(bu_vals[j], 2))
                c.number_format = MONEY
                c.font = font(True)
            for key in keys:
                r = row_of[key]
                users = split.users[key]
                ws.cell(r, col, users).font = font()
                ws.cell(r, col + 1, users / split.total_users).number_format = "0.0%"
                for j in range(3):
                    c = ws.cell(r, col + 2 + j, split.rows[key][j])
                    c.number_format = MONEY
                    c.font = font()
        sums = [sum(split.rows[k][j] for k in data.order) for j in range(3)]
        ws.cell(total_row, col, split.total_users).font = font(True)
        ws.cell(total_row, col + 1, 1).number_format = "0.0%"
        ws.cell(total_row, col + 1).font = font(True)
        for j in range(3):
            c = ws.cell(total_row, col + 2 + j, round(sums[j], 2))
            c.number_format = MONEY
            c.font = font(True)
        for k in range(span):
            ws.cell(total_row, col + k).fill = TOTAL_FILL
        matches = all(abs(sums[j] - (split.excl, split.vat, split.incl)[j]) < 0.005 for j in range(3))
        ws.cell(total_row + 1, col + 2, "Matches the invoice total(s)" if matches else "DOES NOT MATCH the invoice total(s)").font = font(True, brand.GREEN if matches else "9C0006")
        for k in range(span):
            ws.column_dimensions[get_column_letter(col + k)].width = [8, 8, 14, 12, 14][k]
        col += span
    for c in range(1, 3):
        ws.cell(total_row, c).fill = TOTAL_FILL
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 32
    ws.row_dimensions[3].height = 30
    ws.freeze_panes = "C6"


def write_invoices_sheet(wb, invoices, title, name):
    ws = summary.new_sheet(wb, name)
    brand.add_title(ws, f"Invoices used - {title}", "F1")
    heads = ["File", "Supplier", "Subscription", "Invoice number", "Invoice date", "Currency", "Excl VAT", "VAT", "Incl VAT", "Used", "Notes"]
    for i, h in enumerate(heads, 1):
        head(ws.cell(3, i, h), h)
    row = 4
    for inv in invoices:
        values = [inv.file, inv.supplier, inv.subscription, inv.number, inv.date.strftime("%d/%m/%Y") if inv.date else inv.date_text,
                  inv.currency, inv.excl, inv.vat, inv.incl, "Yes" if inv.counted else "NO - dated outside the month", " ".join(inv.notes)]
        for i, v in enumerate(values, 1):
            c = ws.cell(row, i, v)
            c.font = font(not inv.counted and i == 10, "9C0006" if not inv.counted and i == 10 else brand.GREY_DARK)
            if 7 <= i <= 9:
                c.number_format = MONEY
        row += 1
    row += 1
    totals = collections.OrderedDict()
    for inv in invoices:
        if not inv.counted:
            continue
        t = totals.setdefault((inv.subscription, inv.currency), [0.0, 0.0, 0.0])
        t[0] += inv.excl
        t[1] += inv.vat
        t[2] += inv.incl
    ws.cell(row, 1, "Totals per subscription").font = font(True)
    row += 1
    for (subscription, currency), t in totals.items():
        ws.cell(row, 3, subscription).font = font(True)
        ws.cell(row, 6, currency).font = font(True)
        for i, v in enumerate(t):
            c = ws.cell(row, 7 + i, round(v, 2))
            c.number_format = MONEY
            c.font = font(True)
        row += 1
    for letter, width in zip("ABCDEFGHIJK", (44, 12, 20, 22, 13, 10, 14, 12, 14, 28, 80)):
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A4"


def money(value):
    return f"{value:,.2f}"


def split_summary_text(splits):
    lines = []
    for split in splits:
        count = len(split.invoices)
        if split.error:
            lines.append(f"  {block_title(split)}: {split.error}")
            continue
        lines.append(f"  {block_title(split)}: Excl {money(split.excl)} + VAT {money(split.vat)} = Incl {money(split.incl)}   ({split.total_users} users)")
    return "\n".join(lines)


def split_table_text(data, splits):
    good = [s for s in splits if not s.error]
    if not good:
        return "(no invoice could be split)"
    width, cw = 30, 15
    merged = collections.OrderedDict()
    for s in good:
        entry = merged.setdefault((s.subscription, s.currency), {"rows": collections.defaultdict(float), "incl": 0.0})
        for key, values in s.rows.items():
            entry["rows"][key] += values[2]
        entry["incl"] += s.incl
    columns = list(merged.items())
    heads = [(sub if len(sub) <= 10 else sub[:9] + ".") + f" {cur}" for (sub, cur), _ in columns]
    lines = [f"{'Business Unit / Branch':<{width}}" + "".join(f"{h:>{cw}}" for h in heads)]
    lines.append("-" * (width + cw * len(columns)))
    for bu, keys in by_business_unit(data).items():
        lines.append(f"{bu:<{width}}" + "".join(f"{money(sum(e['rows'][k] for k in keys)):>{cw}}" for _, e in columns))
        for key in keys:
            lines.append(f"{'   ' + key[1]:<{width}}" + "".join(f"{money(e['rows'][key]):>{cw}}" for _, e in columns))
        lines.append("")
    lines.append(f"{'TOTAL':<{width}}" + "".join(f"{money(e['incl']):>{cw}}" for _, e in columns))
    return "\n".join(lines)


def split_blocks(splits):
    out = []
    for s in splits:
        out.append({"title": block_title(s), "users": s.total_users, "excl": s.excl, "vat": s.vat, "incl": s.incl,
                    "currency": s.currency, "error": s.error})
    return out


def split_table_data(data, splits):
    good = [s for s in splits if not s.error]
    merged = collections.OrderedDict()
    for s in good:
        entry = merged.setdefault((s.subscription, s.currency), {"rows": collections.defaultdict(float), "incl": 0.0})
        for key, values in s.rows.items():
            entry["rows"][key] += values[2]
        entry["incl"] += s.incl
    columns = list(merged.items())
    heads = [f"{sub} ({cur})" for (sub, cur), _ in columns]
    rows = []
    for bu, keys in by_business_unit(data).items():
        rows.append(("bu", bu, [sum(e["rows"][k] for k in keys) for _, e in columns]))
        for key in keys:
            rows.append(("branch", key[1], [e["rows"][key] for _, e in columns]))
    rows.append(("total", "TOTAL", [e["incl"] for _, e in columns]))
    return heads, rows

import re
import datetime

import engine
from engine import MONTHS, TOTALS_ROWS

IGNORED_REFS = {"amount per user", "inv"}


def number(value):
    return float(value) if isinstance(value, (int, float)) and not isinstance(value, bool) else None


def norm(text):
    return re.sub(r"[^A-Z0-9]", "", str(text).upper()).replace("O", "0")


def section_rows(wb_v, month, year):
    ws = wb_v["MS365 Fees"]
    header = None
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if isinstance(v, datetime.datetime) and v.month == month and v.year == year:
            header = r
            break
    if header is None:
        return None, []
    end = header + 40
    for r in range(header + 1, header + 40):
        v = ws.cell(r, 1).value
        if isinstance(v, datetime.datetime):
            end = r
            break
    rows = []
    for r in range(header + 1, end):
        ref = ws.cell(r, 5).value
        if isinstance(ref, str) and ref.strip() and ref.strip().lower() not in IGNORED_REFS:
            rows.append({"row": r, "ref": ref.strip(), "excl": number(ws.cell(r, 6).value),
                         "vat": number(ws.cell(r, 7).value), "incl": number(ws.cell(r, 8).value)})
    return header, rows


MICROSOFT_NUMBER = re.compile(r"\b([GE][0-9A-Z]{9,10})\b")


def microsoft_number(ref):
    found = MICROSOFT_NUMBER.search(str(ref).upper())
    return norm(found.group(1)) if found else None


def duplicate_rows(wb_v, month, year):
    header, rows = section_rows(wb_v, month, year)
    out = []
    if header is None:
        return out
    earlier = {}
    for m in range(1, month):
        _, previous = section_rows(wb_v, m, year)
        for row in previous:
            number_ = microsoft_number(row["ref"])
            if number_ and number_ not in earlier:
                earlier[number_] = m
    for row in rows:
        number_ = microsoft_number(row["ref"])
        if number_ and number_ in earlier:
            out.append(dict(row, number=number_, earlier_month=earlier[number_]))
    return out


def duplicate_messages(duplicates, month, year):
    out = []
    for d in duplicates:
        amounts = ", ".join(f"{label} {value:,.2f}" for label, value in (("excl", d["excl"]), ("VAT", d["vat"]), ("total", d["incl"])) if value is not None)
        out.append(f"EXCLUDED from {MONTHS[month - 1]}: row {d['row']} of 'MS365 Fees' ({d['ref']}) is a Microsoft invoice that is already in the "
                   f"{MONTHS[d['earlier_month'] - 1]} section, so it was not counted again" + (f" (typed {amounts}; the amounts are kept in a note on that cell)." if amounts else "."))
    return out


def match_row(inv, rows):
    for row in rows:
        if inv.supplier == "Vodacom" and "VB030145" in row["ref"].upper():
            return row
        if inv.number and norm(inv.number) in norm(row["ref"]):
            return row
    return None


def invoice_checks(wb_v, invoices, data, month, year, skip_rows=()):
    out = []
    if not invoices:
        return out
    month_name = MONTHS[month - 1]
    microsoft = [i for i in invoices if i.subscription == "Microsoft 365"]
    if microsoft:
        header, rows = section_rows(wb_v, month, year)
        rows = [r for r in rows if r["row"] not in skip_rows]
        if header is not None:
            used = set()
            differences = []
            for inv in microsoft:
                row = match_row(inv, rows)
                if row is None:
                    out.append(f"{inv.file}: invoice {inv.number} is not listed in the {month_name} section of 'MS365 Fees'.")
                    continue
                used.add(row["row"])
                if inv.currency == "ZAR" and row["incl"] is not None and abs(row["incl"] - inv.incl) > 0.05:
                    percent = (row["incl"] - inv.incl) / inv.incl * 100
                    differences.append(f"{inv.number} (typed {row['incl']:,.2f}, invoice {inv.incl:,.2f}, {percent:+.1f}%)")
            missing = [r for r in rows if r["row"] not in used]
            if missing:
                total = sum(r["excl"] or 0 for r in missing)
                refs = ", ".join(r["ref"] for r in missing[:8])
                out.append(f"'MS365 Fees' lists {len(missing)} {month_name} invoice(s) that have no PDF uploaded ({refs}; R{total:,.2f} excl VAT), "
                           "so they are not in the Microsoft 365 split.")
            if differences:
                out.append("Amounts typed on 'MS365 Fees' differ from the invoice (the split uses the invoice amounts): " + "; ".join(differences[:8]) + ".")
    for inv in invoices:
        if inv.supplier == "Adobe" and inv.lines:
            have = {"pro": data.product_total("adobe acrobat pro"), "std": data.product_total("adobe acrobat standard dc"),
                    "cc": data.product_total("adobe creative suite")}
            billed = {k: inv.lines.get(k, 0) for k in have}
            if billed != have:
                out.append(f"{inv.file}: Adobe bills {billed['pro']} Acrobat Pro, {billed['std']} Standard and {billed['cc']} Creative Cloud, "
                           f"but the user list has {have['pro']}, {have['std']} and {have['cc']}.")
    openai = [i for i in invoices if i.supplier == "OpenAI" and (i.standard_seats or i.premium_seats)]
    if openai:
        latest = max((i for i in openai if i.standard_seats), key=lambda i: i.date or datetime.date.min, default=None)
        premium = sum(i.premium_seats for i in openai)
        have_std, have_prem = data.product_total("chat gpt"), data.product_total("chat gpt premium")
        if latest and latest.standard_seats != have_std:
            out.append(f"OpenAI invoices end at {latest.standard_seats} standard ChatGPT seats (invoice of {latest.date.strftime('%d/%m/%Y') if latest.date else '?'}), "
                       f"but the user list has {have_std}.")
        if premium != have_prem:
            out.append(f"OpenAI invoices bill {premium} Premium seat(s) this month, but the user list has {have_prem}.")
    return out


def formula_notes(wb_f, wb_v, res, plan, month, year, excluded_excl=0.0):
    out = []
    try:
        header, rows = section_rows(wb_v, month, year)
        if header is None or not plan or not plan.pos:
            return out
        fees_f, fees_v = wb_f["MS365 Fees"], wb_v["MS365 Fees"]
        rate_row = next((r for r in range(header + 1, header + 40) if str(fees_v.cell(r, 5).value).strip().upper() == "AMOUNT PER USER"), None)
        if rate_row is None:
            return out
        expense = number(fees_v.cell(rate_row - 1, 6).value)
        if expense:
            expense -= excluded_excl
        formula = fees_f.cell(rate_row, 6).value
        seats = res.seats_total
        if not expense or not isinstance(formula, str) or seats <= 0:
            return out
        divisor_match = re.search(r"/\s*(\d+)\s*\)?\s*$", formula)
        rate_now = None
        if divisor_match:
            divisor = int(divisor_match.group(1))
            rate_now = expense / divisor
            if divisor != int(seats):
                out.append(f"'MS365 Fees' calculates the per-user amount by dividing the month's invoices (R{expense:,.2f}) by {divisor}, "
                           f"but the user list has {int(seats)} seats: R{rate_now:,.2f} per user instead of R{expense / seats:,.2f}. "
                           "This formula was NOT changed because it sets what each entity is charged. Please let the person responsible decide.")
        cc = plan.pos["count_col"]
        cost_col = cc + 1
        letter = engine.get_column_letter(cost_col) if hasattr(engine, "get_column_letter") else None
        if letter is None:
            from openpyxl.utils import get_column_letter
            letter = get_column_letter(cost_col)
        total_formula = wb_f["MS365 Totals"].cell(18, cost_col).value
        if rate_now and isinstance(total_formula, str) and total_formula.replace(" ", "").upper() == f"=SUM({letter}5:{letter}17)":
            costs = {ent: res.entity_seats.get(ent, 0) * rate_now for ent in TOTALS_ROWS if ent != "GWP"}
            sma_group = sum(costs[e] for e in ("SMA", "KAAP", "DBN", "ONLINE", "LTD", "JHB", "WENPRO", "DWF PTA"))
            correct = sma_group + costs["Zire DC"] + costs["VS"]
            shown = sma_group * 2 + costs["Zire DC"] + costs["VS"]
            out.append(f"CORRECTED: on 'MS365 Totals' the USERS cost ({letter}18) added the subtotal in {letter}15 twice and would have shown R{shown:,.2f}. "
                       f"In the new workbook it is R{correct:,.2f}, and VAT and Total ({letter}19 and {letter}20) follow. Earlier months are untouched.")
    except Exception:
        return out
    return out

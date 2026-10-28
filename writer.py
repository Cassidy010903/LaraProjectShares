from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

import engine
from engine import MONTHS, TOTALS_ROWS, ENTITY_NAMES

BOLD = Font(bold=True)
HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")


def put(ws, row, col, value):
    ws.cell(row, col).value = value


def write_totals(res, plan, wb):
    ws = wb["MS365 Totals"]
    pos = plan.pos
    cc = pos["count_col"]
    sma_l = get_column_letter(pos["sma_col"])
    vs_l = get_column_letter(pos["vs_col"])
    for ent, row in TOTALS_ROWS.items():
        put(ws, row, cc, int(res.entity_seats.get(ent, 0)))
    for counts, labels, col, first, last, sum_row in (
            (res.sma_table, plan.sma_labels, pos["sma_col"], engine.SMA_TABLE_FIRST, engine.SMA_TABLE_LAST, engine.SMA_TABLE_SUM),
            (res.vs_table, plan.vs_labels, pos["vs_col"], engine.VS_TABLE_FIRST, engine.VS_TABLE_LAST, engine.VS_TABLE_SUM)):
        by_upper = {k.upper(): v for k, v in counts.items()}
        for label, row in labels.items():
            n = int(by_upper.get(label, 0))
            put(ws, row, col, n if n else None)
        letter = get_column_letter(col)
        put(ws, sum_row, col, f"=SUM({letter}{first}:{letter}{last})")
    cl = get_column_letter(cc)
    put(ws, 15, cc, f"=SUM({cl}5:{cl}14)")
    put(ws, 18, cc, f"=SUM({cl}15:{cl}17)")
    put(ws, 23, cc + 1, int(res.seats_total))
    cost_letter = get_column_letter(cc + 1)
    current = ws.cell(18, cc + 1).value
    if isinstance(current, str) and current.replace(" ", "").upper() == f"=SUM({cost_letter}5:{cost_letter}17)":
        put(ws, 18, cc + 1, f"=SUM({cost_letter}15:{cost_letter}17)")


def write_blocks(res, plan, wb):
    ws = wb["2026"]
    a0, r0, g0 = plan.pos["adobe"], plan.pos["read"], plan.pos["gpt"]
    if a0 is None or r0 is None or g0 is None:
        return
    g = res.groups
    put(ws, 7, a0 + 1, sum(len(g[x]["pro"]) for x in g))
    put(ws, 8, a0 + 1, sum(len(g[x]["std"]) for x in g))
    put(ws, 11, a0 + 1, sum(len(g[x]["cc"]) for x in g))
    put(ws, 15, a0 + 1, len(g["VS"]["pro"]))
    put(ws, 16, a0 + 1, len(g["VS"]["std"]))
    put(ws, 19, a0 + 1, len(g["SMA"]["pro"]))
    put(ws, 20, a0 + 1, len(g["SMA"]["std"]))
    put(ws, 21, a0 + 1, len(g["SMA"]["cc"]))
    put(ws, 23, a0 + 1, len(g["BPK"]["pro"]))
    put(ws, 26, a0 + 1, len(g["BPK"]["std"]))
    put(ws, 27, a0 + 1, len(g["BPK"]["cc"]))
    for c0, key in ((r0, "read"), (g0, "gpt")):
        put(ws, 7, c0 + 1, sum(len(g[x][key]) for x in g))
        for row, grp in ((11, "SMA"), (13, "VS"), (15, "BPK")):
            n = len(g[grp][key])
            if grp == "VS" and not n and ws.cell(row, c0 + 1).value is None:
                continue
            put(ws, row, c0 + 1, n)
            put(ws, row, c0 + 2, engine.join_names(g[grp][key]) if n else None)
            if grp == "VS" and n:
                put(ws, row, c0, "VS")


def write_checks(res, plan, wb, warnings):
    title = f"Checks {MONTHS[res.month - 1][:3]} {res.year}"
    if title in wb.sheetnames:
        del wb[title]
    ws = wb.create_sheet(title)
    put(ws, 1, 1, f"Checks - {MONTHS[res.month - 1]} {res.year}")
    ws["A1"].font = Font(bold=True, size=14)
    row = 3
    put(ws, row, 1, "Users with licences read from the list")
    put(ws, row, 2, res.users_read)
    put(ws, row + 1, 1, "Total Microsoft seats")
    put(ws, row + 1, 2, int(res.seats_total))
    row += 3
    for heading, items in (("Please check", warnings), ("Notes", res.info)):
        put(ws, row, 1, heading)
        ws.cell(row, 1).font = BOLD
        row += 1
        for item in items or ["None"]:
            put(ws, row, 1, item)
            row += 1
        row += 1
    put(ws, row, 1, "Seats per entity")
    ws.cell(row, 1).font = BOLD
    for i, h in enumerate(("Entity", "Last month", "This month"), 1):
        c = ws.cell(row + 1, i, h)
        c.font = BOLD
        c.fill = HEAD_FILL
    row += 2
    for ent in TOTALS_ROWS:
        put(ws, row, 1, ENTITY_NAMES[ent])
        put(ws, row, 2, plan.previous["entities"].get(ent, "n/a") if plan.previous["entities"] else "n/a")
        put(ws, row, 3, int(res.entity_seats.get(ent, 0)))
        row += 1
    row += 1
    for key, heading, now in (("sma", "SMA branches", res.sma_table), ("vs", "Vleissentraal branches", res.vs_table)):
        put(ws, row, 1, heading)
        ws.cell(row, 1).font = BOLD
        for i, h in enumerate(("Branch", "Last month", "This month"), 1):
            c = ws.cell(row + 1, i, h)
            c.font = BOLD
            c.fill = HEAD_FILL
        row += 2
        labels = plan.sma_labels if key == "sma" else plan.vs_labels
        for label in labels:
            put(ws, row, 1, label)
            put(ws, row, 2, plan.previous[key].get(label, "n/a") if plan.previous["entities"] else "n/a")
            put(ws, row, 3, int(next((v for k, v in now.items() if k.upper() == label), 0)))
            row += 1
        row += 1
    put(ws, row, 1, "Licence totals in the user list")
    ws.cell(row, 1).font = BOLD
    row += 1
    for k, v in res.licence_totals.items():
        put(ws, row, 1, k)
        put(ws, row, 2, int(v))
        row += 1
    ws.column_dimensions["A"].width = 60
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 16


def write_figures(plan, wb, figures):
    ws = wb["2026"]
    for subscription, key in (("Adobe", "adobe"), ("Read AI", "read"), ("Chat GPT", "gpt")):
        figure = figures.get(subscription)
        c0 = plan.pos.get(key)
        if not figure or c0 is None:
            continue
        if figure["date"]:
            put(ws, 2, c0, figure["date"])
        if figure["number"]:
            put(ws, 2, c0 + 3, figure["number"])
        if subscription == "Adobe":
            for price_key, row in (("pro", 7), ("std", 8), ("cc", 11)):
                if price_key in figure.get("prices", {}):
                    put(ws, row, c0 + 3, figure["prices"][price_key])
            continue
        if figure["rand"] is None:
            continue
        put(ws, 3, c0 + 3, figure["rand"])
        if figure["usd"] is not None:
            put(ws, 4, c0 + 3, figure["usd"])
        if figure["estimated"]:
            ws.cell(3, c0 + 3).comment = Comment("ESTIMATE: the dollar invoices were converted at the daily ECB exchange rate. "
                                                 "Replace with the rand amount from the bank statement.", "Sully")


def exclude_duplicates(wb, duplicates):
    ws = wb["MS365 Fees"]
    for d in duplicates:
        shown = ", ".join(f"{label} {value:,.2f}" for label, value in (("excl", d["excl"]), ("VAT", d["vat"]), ("total", d["incl"])) if value is not None)
        for column in (6, 7, 8):
            ws.cell(d["row"], column).value = 0
        ws.cell(d["row"], 5).comment = Comment(
            f"EXCLUDED by the program: invoice {d['number']} is already in the {engine.MONTHS[d['earlier_month'] - 1]} section, so it is not counted again. "
            f"Original typed amounts: {shown or 'not available'}.", "Sully")


def write_workbook(res, plan, wb, out_path, warnings, data=None, invoices=None, splits=None, figures=None, duplicates=None):
    if duplicates:
        exclude_duplicates(wb, duplicates)
    write_totals(res, plan, wb)
    write_blocks(res, plan, wb)
    if figures:
        write_figures(plan, wb, figures)
    write_checks(res, plan, wb, warnings)
    if data is not None:
        import summary
        summary.add_summary_sheets(wb, data, res.month, res.year, invoices, splits)
    wb.calculation.fullCalcOnLoad = True
    wb.save(out_path)
